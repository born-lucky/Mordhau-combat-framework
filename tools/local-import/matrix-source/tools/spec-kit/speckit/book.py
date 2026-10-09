# speckit.book - the in-memory spec workbook a populate run fills, and its xlsx I/O (game-agnostic).
#
# A populate run never types a value in: adapters (speckit.adapters) read the game's data through verified readers
# and call Book.source / entity / set / rules / evidence. Book then renders the sheets (Sources, Entities, Fields,
# T_<entity type>, Rules, Evidence, Blobs), compares them cell by cell with the saved workbook (idempotent: unchanged
# content is not rewritten) and writes it with openpyxl in write-only mode.
#
# IDs (permanent): ENT_<CODE>_<key>, FLD_<CODE>_<NAME>, SRC_<KIND>_<name>, RULE_*, EVD_*; CODE per entity type is
# given by the game (Book(codes=...)). Characters outside [A-Za-z0-9_] become "_".
import collections, hashlib, json, math, pathlib, re, struct

ID_RX = re.compile(r"[^A-Za-z0-9_]")


def sid(*parts):
    """an ID from parts: non [A-Za-z0-9_] characters -> '_', empty parts dropped"""
    return "_".join(ID_RX.sub("_", str(p)) for p in parts if p != "")


# ---- numbers ----------------------------------------------------------------------------------------------------
# A game `float` is usually IEEE binary32 (UE, Unity, Godot's stored floats...). Readers often hand it over as f64,
# either the parse of a shortest f32 decimal (0.675) or the widening of raw f32 bits (0.33000001311302185): the same
# f32. The spec stores that f32 as its shortest round-tripping decimal (<= 9 significant digits), which a 16-digit
# spreadsheet cell holds exactly. Consumers read it as f32. A real f64 is the type `double`.
F32_OVERFLOW = 3.4028235677973366e38		# FLT_MAX + half an ulp: anything below rounds to a finite f32
F32_TYPES = {"float", "float_list", "vec2", "vec3", "color"}


def f32(x):
    return struct.unpack("<f", struct.pack("<f", x))[0]


def f32_canon(v):
    if isinstance(v, list): return [f32_canon(x) for x in v]
    if isinstance(v, bool) or not isinstance(v, (int, float)): return v
    if isinstance(v, float) and not math.isfinite(v): raise SystemExit(f"non-finite float {v}")
    if abs(v) >= F32_OVERFLOW: raise SystemExit(f"float {v} is outside f32")
    t = f32(float(v))
    for d in range(1, 10):
        c = float(f"{t:.{d}g}")
        if abs(c) < F32_OVERFLOW and f32(c) == t:
            return c
    return t


def merge_type(a, b, what):
    """the field type when two values of one field disagree (int + float = float; anything + json = json)"""
    if a == b: return a
    if {a, b} == {"int", "float"}: return "float"
    if "json" in (a, b): return "json"
    if {a, b} <= {"float_list", "int_list"}: return "float_list"
    raise SystemExit(f"{what}: type {a} vs {b}")


def infer(v):
    """a spec type for a plain value"""
    if isinstance(v, bool): return "bool"
    if isinstance(v, int): return "int"
    if isinstance(v, float): return "float"
    if isinstance(v, str): return "string"
    if isinstance(v, list):
        if all(isinstance(x, str) for x in v) and v: return "string_list"
        if all(isinstance(x, (int, float)) and not isinstance(x, bool) for x in v) and v: return "float_list"
    return "json"


# ---- the book -----------------------------------------------------------------------------------------------------
class Book:
    """codes: entity type -> ID code (order = sheet order); units: (entity type, field name) -> (unit, unit source);
    version: the game build every source row names"""

    def __init__(self, codes, units=None, version=""):
        self.codes = dict(codes)
        self.units = units or {}
        self.version = version
        self.entities = {}      # id -> row dict
        self.fields = {}        # id -> row dict
        self.values = collections.defaultdict(dict)   # type -> entity id -> {field id: python value}
        self.ftype = {}
        self.sources = {}
        self.rules = {}
        self.evidence = {}

    def source(self, kind, path, cite=""):
        i = sid("SRC", kind.upper(), pathlib.PurePosixPath(path).name if kind == "pkg" else path)[:120]
        if i in self.sources and self.sources[i]["path"] != path:
            # two packages with one name in different folders: the later one gets a stable path-hash suffix
            i = sid(i[:110], hashlib.sha1(path.encode()).hexdigest()[:8])
            if i in self.sources and self.sources[i]["path"] != path:
                raise SystemExit(f"source id collision {i}: {path} vs {self.sources[i]['path']}")
        self.sources[i] = {"source_id": i, "kind": kind, "path": path, "cite": cite, "version": self.version}
        return i

    def entity(self, etype, key, name, source, parent="", record=""):
        i = sid("ENT", self.codes[etype], key)
        if i in self.entities:
            raise SystemExit(f"duplicate entity {i}")
        self.entities[i] = {"entity_id": i, "entity_type": etype, "name": name, "parent_id": parent,
                            "source_id": source, "record": record}
        self.values[etype][i] = {}
        return i

    def field(self, etype, name, stype, ref_type=""):
        i = sid("FLD", self.codes[etype], name.upper())
        f = self.fields.get(i)
        if f is None:
            unit, us = self.units.get((etype, name), ("", "UNCONFIRMED: no source states a unit"))
            self.fields[i] = {"field_id": i, "entity_type": etype, "name": name, "type": stype, "ref_type": ref_type,
                              "unit": unit, "unit_source": us, "source": "", "description": ""}
        elif f["name"] != name:
            raise SystemExit(f"field id collision {i}: {name} vs {f['name']}")
        elif f["type"] != stype:
            f["type"] = merge_type(f["type"], stype, i)
        return i

    def set(self, etype, eid, name, value, stype, ref_type=""):
        fid = self.field(etype, name, stype, ref_type)
        self.values[etype][eid][fid] = f32_canon(value) if stype in F32_TYPES else value


# ---- sheets -------------------------------------------------------------------------------------------------------
SHEET_COLS = {
    "Sources": ["source_id", "kind", "path", "cite", "version"],
    "Entities": ["entity_id", "entity_type", "name", "parent_id", "source_id", "record"],
    "Fields": ["field_id", "entity_type", "name", "type", "ref_type", "unit", "unit_source", "source", "description", "tiers",
               "tier_note"],
    "Rules": ["rule_id", "rule_type", "name", "description", "params", "field_refs", "source", "implemented_by", "area",
              "test_id", "status", "tiers", "tier_note"],
    "Evidence": ["evidence_id", "target_id", "status", "observation", "source", "version", "procedure", "expected",
                 "observed", "test_id"],
}
BLOB_CHUNK = 30000		# under Excel's 32,767 characters per cell


def cell(v, stype):
    """python value -> workbook literal"""
    if v is None: return None
    if stype in ("bool", "int", "float"):
        if isinstance(v, float) and not math.isfinite(v): raise SystemExit(f"non-finite {v}")
        return v
    if stype == "double":
        return repr(float(v))
    if stype in ("string", "ref"):
        s = str(v)
        return json.dumps(s, ensure_ascii=False) if s == "" or s[0] in '"=' else s
    return json.dumps(v, sort_keys=True, separators=(",", ":"), ensure_ascii=False)


def sheets(b):
    """the Book as {sheet: header}, {sheet: rows}"""
    out = {}
    out["Sources"] = [[s[c] for c in SHEET_COLS["Sources"]] for _, s in sorted(b.sources.items())]
    out["Entities"] = [[e[c] for c in SHEET_COLS["Entities"]] for _, e in sorted(b.entities.items())]
    out["Fields"] = [[f.get(c, "") for c in SHEET_COLS["Fields"]] for _, f in sorted(b.fields.items())]
    out["Rules"] = [[r.get(c, "") for c in SHEET_COLS["Rules"]] for _, r in sorted(b.rules.items())]
    out["Evidence"] = [[e[c] for c in SHEET_COLS["Evidence"]] for _, e in sorted(b.evidence.items())]
    hdr = {k: SHEET_COLS[k] for k in ("Sources", "Entities", "Fields", "Rules", "Evidence")}
    blobs = []

    def fit(c, eid, fid):
        # a longer literal goes to the Blobs sheet in parts; the cell says "@blob:<blob id>"
        if not isinstance(c, str) or len(c) <= BLOB_CHUNK: return c
        bid = sid("BLOB", eid, fid)
        for k in range(0, len(c), BLOB_CHUNK):
            blobs.append([bid, k // BLOB_CHUNK, c[k:k + BLOB_CHUNK]])
        return "@blob:" + bid
    for et in b.codes:
        if not b.values[et]:
            continue			# a kind whose dump is absent this run: no sheet
        fids = sorted(f for f, d in b.fields.items() if d["entity_type"] == et)
        name = "T_" + et
        hdr[name] = ["entity_id"] + fids
        out[name] = [[eid] + [fit(cell(vals.get(f), b.fields[f]["type"]), eid, f) for f in fids]
                     for eid, vals in sorted(b.values[et].items())]
    hdr["Blobs"] = ["blob_id", "part", "text"]
    out["Blobs"] = blobs
    return hdr, out


def read_book(path):
    import openpyxl
    path = pathlib.Path(path)
    if not path.exists(): return {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = ws.iter_rows(values_only=True)
        h = list(next(rows, []))
        while h and h[-1] is None: h.pop()
        # read-only rows may stop at the last non-empty cell or run past the header: square them to the header
        out[ws.title] = (h, [(list(r) + [None] * len(h))[:len(h)] for r in rows if any(x is not None for x in r)])
    wb.close()
    return out


def norm(v):
    if isinstance(v, float): return float("%.16g" % v)		# what openpyxl writes (compat/strings.py safe_string)
    return None if v == "" else v


def diff(old, hdr, new):
    """per-sheet changes between a saved workbook and new sheets (rows keyed by ID)"""
    rep = []
    for name in list(hdr) + [n for n in old if n not in hdr]:
        if name not in hdr: rep.append(f"  {name}: sheet removed"); continue
        if name not in old: rep.append(f"  {name}: new sheet, {len(new[name])} rows"); continue
        oh, orows = old[name]
        if [norm(x) for x in oh] != [norm(x) for x in hdr[name]]:
            rep.append(f"  {name}: columns changed ({len(oh)} -> {len(hdr[name])})")
        key = (lambda r: (r[0], r[1])) if name == "Blobs" else (lambda r: r[0])
        om = {key(r): [norm(x) for x in r] for r in orows}
        nm = {key(r): [norm(x) for x in r] for r in new[name]}
        add = [k for k in nm if k not in om]; rem = [k for k in om if k not in nm]
        chg = [k for k in nm if k in om and om[k] != nm[k]]
        if add or rem or chg:
            rep.append(f"  {name}: +{len(add)} -{len(rem)} ~{len(chg)}" + (f" (e.g. {(add + rem + chg)[:3]})"))
    return rep


def write_book(path, hdr, data):
    import openpyxl
    path = pathlib.Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook(write_only=True)
    for name in hdr:
        ws = wb.create_sheet(name)
        ws.append(hdr[name])
        for r in data[name]:
            ws.append(r)
    tmp = path.with_suffix(".tmp.xlsx")
    wb.save(tmp)
    tmp.replace(path)


def save(b, path, check=False, log=print):
    """render, diff against the saved workbook, write if changed (unless check). Returns 0, or 1 when check finds it
    stale."""
    hdr, data = sheets(b)
    rep = diff(read_book(path), hdr, data)
    if not rep:
        log(f"{path}: up to date ({sum(len(v) for v in data.values())} rows)")
        return 0
    log(f"{path}: {'STALE' if check else 'changed'}:")
    log("\n".join(rep))
    if check:
        return 1
    write_book(path, hdr, data)
    log(f"wrote {path}")
    return 0
